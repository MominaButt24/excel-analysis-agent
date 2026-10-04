import re
from decimal import Decimal, InvalidOperation
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from parser.normalizer import normalize_value
from parser.structure import (
    detect_header_row,
    detect_actual_header_row,
    detect_active_columns,
    detect_parent_headers,
    detect_blocks,
    get_row_values,
    is_empty_row,
)


def parse_row(
    worksheet,
    row_number,
    columns,
    sheet_name,
):
    cells = {}
    values = {}

    for column_number in columns:

        cell = worksheet.cell(
            row=row_number,
            column=column_number
        )

        value = normalize_value(
            cell.value
        )

        column_letter = get_column_letter(
            column_number
        )

        if value is None:
            for merged_range in worksheet.merged_cells.ranges:
                if (
                    merged_range.min_col == column_number
                    and merged_range.max_col == column_number
                    and merged_range.min_row <= row_number <= merged_range.max_row
                ):
                    value = normalize_value(
                        worksheet.cell(
                            row=merged_range.min_row,
                            column=column_number,
                        ).value
                    )
                    break

        values[column_letter] = value

        cell_value = normalize_value(cell.value)
        if cell_value is None:
            continue

        cells[cell.coordinate] = {
            "column": column_number,
            "value": cell_value,
        }

    first_text_value = next(
        (
            str(cell["value"]).strip()
            for cell in cells.values()
            if isinstance(cell["value"], str)
            and cell["value"].strip()
        ),
        "",
    )
    normalized_label = first_text_value.casefold()

    if re.match(
        r"^(?:grand\s+total|sub\s*total|subtotal|total)\b",
        normalized_label,
    ):
        row_kind = "summary"
    elif re.match(r"^note\s*:", normalized_label):
        row_kind = "note"
    else:
        row_kind = "data"

    return {
        "sheet": sheet_name,
        "row": row_number,
        "kind": row_kind,
        "values": values,
        "cells": cells,
    }


def parse_numeric_text(value, allow_embedded=False):
    if not isinstance(value, str):
        return None

    normalized = value.strip().replace(",", "")
    if not normalized:
        return None

    numeric_text = normalized

    try:
        number = Decimal(normalized)
    except InvalidOperation:
        if not allow_embedded:
            return None

        currency_match = re.search(
            r"(?:rs\.?|inr|usd|eur|gbp|\$|\u20b9|\u20ac|\u00a3)\s*"
            r"([+-]?(?:\d+(?:\.\d*)?|\.\d+))",
            value,
            flags=re.IGNORECASE,
        )
        amount_match = re.search(
            r"(?:fine|amount|price|cost|total|salary|payment|charge|fee)"
            r"\D{0,16}([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)"
            r"(?:\.\d+)?)",
            value,
            flags=re.IGNORECASE,
        )
        number_match = currency_match or amount_match
        if number_match is None:
            return None

        matched_text = number_match.group(1)
        numeric_text = matched_text.replace(",", "")
        try:
            number = Decimal(numeric_text)
        except InvalidOperation:
            return None

    if "." in numeric_text or "e" in numeric_text.lower():
        return float(number)

    return int(number)


def infer_numeric_columns(rows, columns):
    numeric_columns = set()
    data_rows = [
        row
        for row in rows
        if row.get("kind", "data") == "data"
    ]

    for column in columns:
        field_id = column["field_id"]
        header = str(column.get("name") or "").casefold()
        is_amount_column = any(
            term in header
            for term in (
                "fine",
                "amount",
                "price",
                "cost",
                "total",
                "salary",
                "payment",
                "charge",
                "fee",
            )
        )

        populated_values = [
            row["values"][field_id]
            for row in data_rows
            if row["values"][field_id] is not None
        ]

        native_numeric_count = sum(
            isinstance(value, (int, float))
            and not isinstance(value, bool)
            for value in populated_values
        )

        if not populated_values or (
            not native_numeric_count
            and not is_amount_column
        ):
            continue

        numeric_candidates = []
        numeric_count = 0

        for row in data_rows:
            value = row["values"][field_id]
            if value is None:
                continue

            converted = parse_numeric_text(
                value,
                allow_embedded=is_amount_column,
            )
            numeric_candidates.append((row, converted))

            if (
                isinstance(value, (int, float))
                and not isinstance(value, bool)
            ) or converted is not None:
                numeric_count += 1

        if numeric_count / len(populated_values) < 0.8:
            continue

        numeric_columns.add(field_id)
        for row, converted in numeric_candidates:
            if converted is not None:
                row["values"][field_id] = converted

    for column in columns:
        column["inferred_type"] = (
            "number"
            if column["field_id"] in numeric_columns
            else "unknown"
        )


def classify_indexed_date_rows(rows, columns):
    date_columns = [
        column
        for column in columns
        if "date" in str(column.get("name") or "").casefold()
    ]
    index_columns = [
        column
        for column in columns
        if re.fullmatch(
            r"(?:sr\.?\s*no\.?|serial(?:\s*(?:no\.?|number))?|index)",
            str(column.get("name") or "").strip().casefold(),
        )
    ]

    if not date_columns or not index_columns:
        return

    date_fields = [column["field_id"] for column in date_columns]
    index_fields = [column["field_id"] for column in index_columns]

    for row in rows:
        if row.get("kind", "data") != "data":
            continue

        has_record_date = any(
            isinstance(row["values"].get(field_id), str)
            and _is_iso_date(row["values"][field_id])
            for field_id in date_fields
        )
        has_record_index = any(
            isinstance(row["values"].get(field_id), (int, float))
            and not isinstance(row["values"].get(field_id), bool)
            for field_id in index_fields
        )

        if has_record_date and has_record_index:
            continue

        populated_values = [
            value
            for value in row["values"].values()
            if value is not None
        ]

        if (
            len(populated_values) == 1
            and any(
                field_id in index_fields
                and isinstance(row["values"].get(field_id), str)
                and _is_iso_date(row["values"][field_id])
                for field_id in index_fields
            )
        ):
            row["kind"] = "section_header"
            continue

        is_leave_marker = any(
            isinstance(value, str)
            and value.strip().casefold() in {"leave", "leave day"}
            for value in populated_values
        )
        if is_leave_marker and not has_record_date and not has_record_index:
            row["kind"] = "non_record"


def _is_iso_date(value):
    try:
        date.fromisoformat(value)
        return True
    except (TypeError, ValueError):
        return False


def build_columns(
    worksheet,
    header_row,
    active_columns,
):
    """
    Build actual child columns.

    For normal sheets:
        header_row = actual header row.

    For grouped sheets:
        header_row = child header row.
    """

    columns = []

    for column_number in active_columns:

        cell = worksheet.cell(
            row=header_row,
            column=column_number
        )

        name = normalize_value(
            cell.value
        )

        columns.append({
            "index": column_number,
            "field_id": get_column_letter(column_number),
            "name": name,
            "coordinate": cell.coordinate,
        })

    return columns


def fill_merged_header_names(
    worksheet,
    columns,
    parent_headers,
):
    """
    Fill child headers represented by merged cells.

    Handles:
        Horizontal:
            A4:C4 = Item Details

        Vertical:
            H4:H5 = Status
    """

    for column in columns:

        if column["name"] is not None:
            continue

        column_number = column["index"]

        header_row = worksheet[
            column["coordinate"]
        ].row

        for merged_range in worksheet.merged_cells.ranges:

            if column_number not in range(
                merged_range.min_col,
                merged_range.max_col + 1
            ):
                continue

            if header_row not in range(
                merged_range.min_row,
                merged_range.max_row + 1
            ):
                continue

            value = worksheet.cell(
                row=merged_range.min_row,
                column=merged_range.min_col
            ).value

            if value is not None:
                column["name"] = str(value).strip()
                break

    return columns


def parse_sheet(worksheet):

    detected_header_row = detect_header_row(
        worksheet
    )

    actual_header_row = detect_actual_header_row(
        worksheet,
        detected_header_row
    )

    active_columns = detect_active_columns(
        worksheet,
        start_row=actual_header_row or 1,
    )

    parent_headers = detect_parent_headers(
        worksheet,
        detected_header_row,
        actual_header_row,
    )

    columns = []

    if actual_header_row:

        columns = build_columns(
            worksheet,
            actual_header_row,
            active_columns,
        )

        columns = fill_merged_header_names(
            worksheet,
            columns,
            parent_headers,
        )

    blocks = []

    if actual_header_row:

        detected_blocks = detect_blocks(
            worksheet,
            actual_header_row + 1,
        )

        for block_index, row_numbers in enumerate(
            detected_blocks,
            start=1
        ):

            rows = []

            for row_number in row_numbers:

                row = parse_row(
                    worksheet,
                    row_number,
                    active_columns,
                    worksheet.title,
                )

                if row["cells"]:
                    rows.append(row)

            if rows:

                blocks.append({
                    "block_id": block_index,
                    "rows": rows,
                })

    else:

        rows = []

        for row_number in range(
            1,
            worksheet.max_row + 1
        ):

            values = get_row_values(
                worksheet,
                row_number
            )

            if is_empty_row(values):
                continue

            row = parse_row(
                worksheet,
                row_number,
                active_columns,
                worksheet.title,
            )

            rows.append(row)

        if rows:

            blocks.append({
                "block_id": 1,
                "rows": rows,
            })

    all_rows = [
        row
        for block in blocks
        for row in block["rows"]
    ]
    infer_numeric_columns(all_rows, columns)
    classify_indexed_date_rows(all_rows, columns)

    return {
        "name": worksheet.title,

        "dimensions": {
            "rows": worksheet.max_row,
            "columns": worksheet.max_column,
        },

        # Keep the original detected row for
        # backwards compatibility.
        "header_row": detected_header_row,

        # New semantic header information.
        "actual_header_row": actual_header_row,

        "parent_headers": parent_headers,

        "columns": columns,

        "merged_ranges": [
            str(r)
            for r in worksheet.merged_cells.ranges
        ],

        "blocks": blocks,
    }


def parse_xlsx(file_path):

    workbook = load_workbook(
        file_path,
        data_only=True,
    )

    return {
        "file": file_path,

        "sheets": [
            parse_sheet(worksheet)
            for worksheet in workbook.worksheets
        ],
    }