import sys
import json
from typing import Any

# ------------------------------------------------------------------
# Project source path inside the sandbox
# ------------------------------------------------------------------

sys.path.insert(0, "/workspace/excel_agent_src")


from openpyxl import load_workbook

from parser.xlsx_parser import parse_xlsx
from grounding.metadata import build_workbook_metadata
from executor.executor import execute_plan
from executor.plan import ExecutionPlan


# ------------------------------------------------------------------
# Configuration
# ------------------------------------------------------------------

WORKBOOK = "/workspace/workbook.xlsx"

MAX_ROWS = 20
MAX_MATCHES = 20


# ------------------------------------------------------------------
# Logging
# ------------------------------------------------------------------

def log(message: str):
    """
    Print execution logs to stderr.

    IMPORTANT:
    stdout is reserved for the final JSON result because the
    parent tool parses stdout with json.loads().
    """
    print(
        f"[SANDBOX] {message}",
        file=sys.stderr,
        flush=True,
    )


# ------------------------------------------------------------------
# Workbook loading
# ------------------------------------------------------------------

def load_workbook_data():
    """
    Parse the workbook using the project's deterministic parser
    and build structural metadata.
    """

    log("Loading workbook through XLSX parser...")

    parsed = parse_xlsx(WORKBOOK)

    log("Building workbook metadata...")

    metadata = build_workbook_metadata(parsed)

    log("Workbook loaded successfully.")

    return parsed, metadata


# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------

def make_json_safe(value: Any):
    """
    Convert values such as dates into JSON-safe values.
    """

    if isinstance(value, dict):
        return {
            str(key): make_json_safe(item)
            for key, item in value.items()
        }

    if isinstance(value, list):
        return [
            make_json_safe(item)
            for item in value
        ]

    if hasattr(value, "isoformat"):
        return value.isoformat()

    return value


def get_sheet_metadata(
    metadata: dict,
    sheet_name: str,
):
    """
    Resolve a sheet name case-insensitively.
    """

    for sheet in metadata["sheets"]:
        if sheet["name"].lower() == sheet_name.lower():
            return sheet

    return None


def compact_sheet_metadata(sheet: dict):
    """
    Return only the structural metadata the agent needs.
    """

    return {
        "name": sheet["name"],
        "dimensions": sheet["dimensions"],
        "header_row": sheet["header_row"],
        "columns": sheet["columns"],
        "blocks": [
            {
                "block_id": block["block_id"],
                "row_count": len(block["rows"]),
                "start_row": min(block["rows"])
                if block["rows"]
                else None,
                "end_row": max(block["rows"])
                if block["rows"]
                else None,
            }
            for block in sheet["blocks"]
        ],
    }


# ------------------------------------------------------------------
# Workbook inspection
# ------------------------------------------------------------------

def inspect_workbook():
    """
    Inspect workbook structure.

    Returns:
    - sheets
    - dimensions
    - header rows
    - columns
    - structural blocks

    This does NOT dump workbook contents.
    """

    log("Inspecting workbook structure...")

    _, metadata = load_workbook_data()

    result = {
        "sheets": [
            compact_sheet_metadata(sheet)
            for sheet in metadata["sheets"]
        ]
    }

    log(
        f"Workbook inspection complete: "
        f"{len(result['sheets'])} sheets found."
    )

    return result


def inspect_sheet(sheet_name: str):
    """
    Inspect one worksheet's structure.
    """

    log(f"Inspecting sheet: {sheet_name}")

    _, metadata = load_workbook_data()

    sheet = get_sheet_metadata(
        metadata,
        sheet_name,
    )

    if sheet is None:
        log(f"Sheet not found: {sheet_name}")

        return {
            "error": f"Sheet not found: {sheet_name}"
        }

    result = compact_sheet_metadata(sheet)

    log(
        f"Sheet inspection complete: "
        f"{sheet['name']}"
    )

    return result


# ------------------------------------------------------------------
# Deterministic Excel operations
# ------------------------------------------------------------------

def excel_query(
    sheet: str,
    operation: str,
    column: str | None = None,
    filters: list[dict] | None = None,
    query: str | None = None,
):
    """
    Execute one of the deterministic Excel operations
    implemented by the project's executor.

    The Deep Agent decides WHAT operation is needed.
    The project's executor performs the actual calculation.
    """

    log(
        f"Excel query: "
        f"sheet={sheet}, "
        f"operation={operation}, "
        f"column={column}"
    )

    parsed_workbook, _ = load_workbook_data()

    # --------------------------------------------------------------
    # Normalize filters at the tool boundary.
    # --------------------------------------------------------------

    normalized_filters = []

    for condition in filters or []:

        condition = dict(condition)

        if condition.get("operator") == "=":
            condition["operator"] = "=="

        if (
            condition.get("column") is None
            and column is not None
        ):
            condition["column"] = column

        normalized_filters.append(condition)

    # --------------------------------------------------------------
    # Generic "contains" → existing search capability
    # --------------------------------------------------------------

    contains_filters = [
        condition
        for condition in normalized_filters
        if condition.get("operator") == "contains"
    ]

    if (
        operation == "filter"
        and len(contains_filters) == 1
    ):
        condition = contains_filters[0]

        column = condition.get(
            "column",
            column,
        )

        query = condition.get(
            "value",
            query,
        )

        operation = "search"
        normalized_filters = []

        log(
            "Translated single 'contains' filter "
            "into deterministic search operation."
        )

    # --------------------------------------------------------------
    # Build deterministic execution plan
    # --------------------------------------------------------------

    plan = ExecutionPlan(
        sheet=sheet,
        operation=operation,
        column=column,
        filters=normalized_filters,
        query=query,
    )

    log("Executing deterministic Excel plan...")

    result = execute_plan(
        parsed_workbook,
        plan,
    )

    result = make_json_safe(result)

    # --------------------------------------------------------------
    # Bound list results
    # --------------------------------------------------------------

    if isinstance(result, list):

        total = len(result)

        limited = result[:MAX_ROWS]

        response = {
            "result": limited,
            "returned_rows": len(limited),
            "total_rows": total,
            "truncated": total > MAX_ROWS,
        }

        log(
            f"Excel query complete: "
            f"{len(limited)}/{total} rows returned."
        )

        return response

    response = {
        "result": result,
        "truncated": False,
    }

    log("Excel query complete.")

    return response


# ------------------------------------------------------------------
# Raw workbook text search
# ------------------------------------------------------------------

def search_workbook_text(
    query: str,
    sheet_name: str | None = None,
):
    """
    Search actual workbook cells.

    This is intentionally separate from the deterministic
    executor because some useful information exists outside
    normal structured data rows:

    - leave markers
    - notes
    - section headers
    - merged labels
    - names
    - project names
    - other raw workbook text
    """

    log(
        f"Searching workbook text: "
        f"query={query!r}, "
        f"sheet={sheet_name}"
    )

    wb = load_workbook(
        WORKBOOK,
        data_only=True,
    )

    # --------------------------------------------------------------
    # Resolve optional sheet case-insensitively
    # --------------------------------------------------------------

    if sheet_name:

        target_sheet = None

        for ws in wb.worksheets:
            if ws.title.lower() == sheet_name.lower():
                target_sheet = ws
                break

        if target_sheet is None:
            return {
                "matches": [],
                "returned_matches": 0,
                "truncated": False,
                "error": f"Sheet not found: {sheet_name}",
            }

        sheets = [target_sheet]

    else:
        sheets = wb.worksheets

    # --------------------------------------------------------------
    # Search
    # --------------------------------------------------------------

    query_lower = query.lower()

    matches = []

    for ws in sheets:

        for row in ws.iter_rows():

            for cell in row:

                if cell.value is None:
                    continue

                if query_lower not in str(
                    cell.value
                ).lower():
                    continue

                row_values = {}

                for row_cell in ws[cell.row]:

                    if row_cell.value is not None:
                        row_values[
                            row_cell.coordinate
                        ] = make_json_safe(
                            row_cell.value
                        )

                matches.append(
                    {
                        "sheet": ws.title,
                        "coordinate": cell.coordinate,
                        "value": make_json_safe(
                            cell.value
                        ),
                        "row": cell.row,
                        "row_values": row_values,
                    }
                )

                if len(matches) >= MAX_MATCHES:

                    log(
                        f"Text search reached "
                        f"MAX_MATCHES={MAX_MATCHES}."
                    )

                    return {
                        "matches": matches,
                        "returned_matches": len(matches),
                        "truncated": True,
                    }

    log(
        f"Text search complete: "
        f"{len(matches)} matches found."
    )

    return {
        "matches": matches,
        "returned_matches": len(matches),
        "truncated": False,
    }


# ------------------------------------------------------------------
# CLI entry point
# ------------------------------------------------------------------

if __name__ == "__main__":

    log("Workbook runner started.")

    result = inspect_workbook()

    # IMPORTANT:
    # Only JSON goes to stdout.
    # Logs go to stderr.
    print(
        json.dumps(
            result,
            default=str,
        )
    )
